# -*- coding: utf-8 -*-
"""Заполняет пропущенные координаты в «Локации» (Nominatim + приблизительные для диких мест)."""
import openpyxl

XLSX = "Sevastopol AI База.xlsx"

COORDS = {
    # точные (OSM/Nominatim)
    "Приморский бульвар": "44.617031, 33.523836",
    "Херсонес": "44.610846, 33.490369",
    "Ласпинская смотровая": "44.430924, 33.714012",
    "Байдарские ворота": "44.406184, 33.782023",
    "Сапун-гора": "44.555924, 33.585154",
    "Грот Дианы": "44.509850, 33.478403",
    "Кадыковский карьер": "44.516074, 33.568122",
    # приблизительные (в OSM объекта нет, точка района/берега)
    "Затерянный мир": "44.463600, 33.642700",
    "Скала Святого Явления": "44.502900, 33.505800",
}

wb = openpyxl.load_workbook(XLSX)
ws = wb["Локации"]
filled = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 2).value
    if name and str(name).strip() in COORDS and not ws.cell(r, 5).value:
        ws.cell(r, 5, COORDS[str(name).strip()])
        filled.append(str(name).strip())
wb.save(XLSX)
print("Заполнено координат в Локации:", len(filled))
for n in filled:
    print(" -", n)

# контроль: дыры в Где поесть
wb = openpyxl.load_workbook(XLSX)
ws = wb["Где поесть"]
holes = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 3).value
    coord = ws.cell(r, 8).value
    if name and not coord:
        holes.append((r, str(name).strip()))
print("\nГде поесть без координат:", holes if holes else "нет")

ws = wb["Локации"]
holes = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 2).value
    coord = ws.cell(r, 5).value
    if name and not coord:
        holes.append((r, str(name).strip()))
print("Локации без координат:", holes if holes else "нет")
