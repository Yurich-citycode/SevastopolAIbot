# -*- coding: utf-8 -*-
import openpyxl

XLSX = "Sevastopol AI База.xlsx"

CONF = [
    ("Медоборы", "✔️ Сеть кондитерских: торты, пирожные, выпечка\n✔️ Торт «Ай-Петри», свежие десерты к чаю", "улица Адмирала Октябрьского, 20", "Центр", "44.602421, 33.516187"),
    ("Медоборы", "✔️ Кондитерская на Октябрьской Революции\n✔️ Торты и пирожные собственного производства", "проспект Октябрьской Революции, 57А", "ПОР", "44.590225, 33.460958"),
]

wb = openpyxl.load_workbook(XLSX)
ws = wb["Где поесть"]
last = 1
for r in range(2, ws.max_row + 1):
    if any(c.value not in (None, "") for c in ws[r]):
        last = r
start = last + 1
for i, (name, desc, addr, district, coord) in enumerate(CONF):
    r = start + i
    ws.cell(r, 1, "Кондитерские")
    ws.cell(r, 3, name)
    ws.cell(r, 4, desc)
    ws.cell(r, 6, addr)
    ws.cell(r, 7, district)
    ws.cell(r, 8, coord)
wb.save(XLSX)
print(f"OK: кондитерских добавлено {len(CONF)}, строки {start}–{start+len(CONF)-1}")
