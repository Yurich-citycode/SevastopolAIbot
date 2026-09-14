# -*- coding: utf-8 -*-
"""Батч 4: три заведения на проспекте Гагарина, геокодированные по номеру дома."""
import openpyxl

XLSX = "Sevastopol AI База.xlsx"

ROWS = [
    ("Кафе", "Бары и пабы", "Пинта",
     "✔️ Пивной бар на проспекте Гагарина\n✔️ Разливное пиво, закуски",
     "проспект Гагарина, 44", "Стрелка", "44.601834, 33.479116", None, None),
    ("Пекарни", None, "Хлебный",
     "✔️ Пекарня на проспекте Гагарина\n✔️ Свежая выпечка и хлеб",
     "проспект Гагарина, 11", "Стрелка", "44.599669, 33.487230", None, None),
    ("Кафе", None, "Али-Баба",
     "✔️ Кафе в ТЦ «Стрелецкий»\n✔️ Восточная кухня, быстро и сытно",
     "проспект Гагарина, 8 (ТЦ «Стрелецкий»)", "Стрелка", "44.599691, 33.489818", None, None),
]

wb = openpyxl.load_workbook(XLSX)
ws = wb["Где поесть"]
last = 1
for r in range(2, ws.max_row + 1):
    if any(c.value not in (None, "") for c in ws[r]):
        last = r
start = last + 1
for i, (cat, sub, name, desc, addr, district, coord, phone, hours) in enumerate(ROWS):
    r = start + i
    ws.cell(r, 1, cat)
    if sub:
        ws.cell(r, 2, sub)
    ws.cell(r, 3, name)
    ws.cell(r, 4, desc)
    ws.cell(r, 6, addr)
    ws.cell(r, 7, district)
    ws.cell(r, 8, coord)
    if phone:
        ws.cell(r, 14, phone)
    if hours:
        ws.cell(r, 15, hours)
wb.save(XLSX)
print(f"OK: батч 4 — добавлено {len(ROWS)} строк, строки {start}–{start+len(ROWS)-1}")
